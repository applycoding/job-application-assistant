#!/usr/bin/env python3
"""Write ``samples/alex-doe-volunteering.docx`` with stable zip metadata."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent.parent / "samples" / "alex-doe-volunteering.docx"
_FIXED_DT = (1980, 1, 1, 0, 0, 0)


def _build_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("FICTIONAL SAMPLE DATA - Alex Doe is not a real person.")
    doc.add_heading("Volunteering - Alex Doe", level=1)
    doc.add_paragraph(
        "Food Bank of Example City (2019–2021): built a simple Postgres + Metabase "
        "dashboard for donation trends."
    )
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Role"
    table.rows[0].cells[1].text = "Data volunteer"
    table.rows[1].cells[0].text = "Hours"
    table.rows[1].cells[1].text = "~4 per month"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _normalize_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data), "r")
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w", compression=zipfile.ZIP_DEFLATED) as dest:
        for info in src.infolist():
            payload = src.read(info.filename)
            new_info = zipfile.ZipInfo(filename=info.filename)
            new_info.compress_type = zipfile.ZIP_DEFLATED
            new_info.date_time = _FIXED_DT
            new_info.external_attr = info.external_attr
            dest.writestr(new_info, payload)
    return out_buf.getvalue()


def main() -> None:
    raw = _build_docx_bytes()
    OUT.write_bytes(_normalize_zip(raw))


if __name__ == "__main__":
    main()
